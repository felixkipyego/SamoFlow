# backend/app/ingest/extract_docx.py
# Task 2.4.e: DOCX extraction primitive -- shared logic (Step 2.4 decision
# (b)'s own precedent), called by Step 2.7's upload adapter, not
# duplicated or rebuilt there. Reuses 2.4.b's own ExtractedContent/
# ContentBlock shape unchanged -- confirmed this is the right shape here:
# unlike PDF's plain text, a DOCX genuinely has semantic heading
# structure (built-in paragraph styles), so it produces HTML's richer
# multi-block shape, not PDF's flattened single block.
#
# heading_path algorithm: the SAME document-order stack algorithm as
# extract_html.py's own (walk paragraphs in order, maintain a stack of
# currently-open headings keyed by level, a new heading pops every
# stacked entry at its own level or deeper then pushes itself) -- but
# simpler to apply here, since python-docx's own `document.paragraphs`
# is ALREADY a flat, ordered sequence with no nesting concept at all.
# HTML needed the "h3 is a DOM SIBLING of h2, never its child" realization
# before the stack approach made sense; DOCX paragraphs were never nested
# to begin with, so that realization isn't needed here -- just applying
# the same already-proven algorithm to an already-flat sequence.
#
# Heading-level detection, confirmed live, not assumed: paragraph.style.
# name for Word's built-in heading styles is exactly "Heading 1" through
# "Heading 9" (confirmed by building and re-reading a real probe
# document) -- Word's own level-0 "Title" style is a SEPARATE named style
# ("Title", not "Heading 0"), confirmed not to collide. Scoped to 1-6
# only, matching HTML's own h1-h6 scope exactly (symmetric design, not
# arbitrary) -- "Heading 7/8/9" exist as real built-in styles but are
# out of scope, same as there being no h7/h8/h9 in HTML. "TOC Heading"
# (the auto-generated table-of-contents style) does NOT match the
# "Heading N" pattern (it ends with, not starts with, "Heading"),
# confirmed not to collide. Only PARAGRAPH-type styles are ever read via
# paragraph.style.name -- the separate "Heading N Char" CHARACTER styles
# are a different style kind entirely and never appear here. Non-English
# Word templates (e.g. "Titre 1") are not handled -- already covered by
# this project's own existing "English only for v1" constraint (docs/
# SPEC.md §3), not a new gap.
#
# title: document.core_properties.title, normalized "" -> None (confirmed
# live: python-docx returns an empty STRING, not None, when no title is
# set -- normalized here so ExtractedContent's own "no title" contract
# stays consistently None, matching extract_html.py's own behavior).
# Same honest caveat as PDF's /Info /Title: may be empty/unreliable for
# a real document depending on how it was authored -- but unlike PDF,
# this field is at least structurally meaningful when populated (a real
# Word "File > Properties > Title" field, not a loosely-defined PDF
# producer convention).
#
# Malformed input: left to propagate UNCAUGHT and unwrapped, same
# reasoning as extract_pdf.py. Confirmed live: a truncated/corrupted DOCX
# (a DOCX is a ZIP archive) raises zipfile.BadZipFile -- a standard
# library exception surfaced through python-docx's own ZIP-based
# parsing, already clear, specific, and documented. python-docx also
# has its own docx.opc.exceptions.PackageNotFoundError for certain other
# package-level failures; neither is caught or re-wrapped here.
import re
from io import BytesIO

from docx import Document

from app.ingest.extract_html import (
    ContentBlock,
    ExtractedContent,
    current_heading_path,
    push_heading,
)

_HEADING_STYLE_RE = re.compile(r"^Heading ([1-6])$")


def _heading_level(style_name: str) -> int | None:
    match = _HEADING_STYLE_RE.match(style_name)
    return int(match.group(1)) if match else None


def extract_docx(docx_bytes: bytes) -> ExtractedContent:
    document = Document(BytesIO(docx_bytes))
    title = document.core_properties.title or None

    blocks: list[ContentBlock] = []
    heading_stack: list[tuple[int, str]] = []
    all_paragraph_texts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        all_paragraph_texts.append(text)

        level = _heading_level(paragraph.style.name)
        if level is not None:
            heading_stack = push_heading(heading_stack, level, text)
            continue

        blocks.append(ContentBlock(heading_path=current_heading_path(heading_stack), text=text))

    word_count = len(" ".join(all_paragraph_texts).split())
    return ExtractedContent(title=title, blocks=tuple(blocks), word_count=word_count)
