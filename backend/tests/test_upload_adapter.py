# backend/tests/test_upload_adapter.py
# Task 2.7.c [SECURITY]: tests for app/ingest/upload_adapter.py's
# ingest_upload() -- live against BOTH real services (test-db AND
# test-qdrant), using live_test_services() (tests/conftest.py), matching
# test_web_adapter.py's own ingest_url() test methodology exactly. No
# HTTP server anywhere in this file -- uploads have no fetch step at
# all; upload_storage.save() under a real pytest tmp_path IS "the
# upload" here.
import io
import uuid
import zipfile

import pytest
import sqlalchemy as sa

from app import qdrant
from app.ingest import upload_adapter as upload_adapter_module
from app.ingest import upload_storage
from app.ingest.models import Document, Source
from app.ingest.upload_adapter import ingest_upload
from app.tenancy.models import Tenant
from tests.conftest import (
    db_session,
    patch_embed_dense,
    read_docx_fixture,
    read_markdown_fixture,
    read_pdf_fixture,
    read_text_fixture,
)

# Settings.docx_max_part_size_bytes's own default (2.7.b) -- a plain test
# constant, not a Settings read, matching ingest_upload()'s own "plain
# parameters, no Settings coupling" design.
_MAX_DOCX_PART_SIZE = 100 * 1024 * 1024
_MAX_UPLOAD_SIZE = 50 * 1024 * 1024


async def _seed_tenant_and_upload_source() -> dict:
    # No verified_domain at all -- uploads have no fetch step, so there
    # is nothing to verify, unlike test_web_adapter.py's own identical-
    # shaped helper for the `urls` adapter.
    tenant_id = uuid.uuid4()
    async with db_session() as session:
        session.add(Tenant(id=tenant_id, name="Upload Adapter Proof Tenant", status="active"))
        await session.commit()

    async with db_session() as session:
        source = Source(
            tenant_id=tenant_id, type="upload", refresh_interval="daily", status="active"
        )
        session.add(source)
        await session.commit()
        source_id = source.id

    return {"tenant_id": tenant_id, "source_id": source_id}


async def _setup(monkeypatch, live_test_services):
    patch_embed_dense(monkeypatch)
    _, client, collection_name = live_test_services
    await qdrant.ensure_collection(client, collection_name)
    ids = await _seed_tenant_and_upload_source()
    return client, collection_name, ids


async def _ingest(
    client,
    collection_name: str,
    *,
    tenant_id,
    source_id,
    upload_id: uuid.UUID,
    original_filename: str,
    storage_dir,
):
    async with db_session() as session:
        result = await ingest_upload(
            session,
            client,
            collection_name,
            tenant_id=tenant_id,
            source_id=source_id,
            upload_id=upload_id,
            original_filename=original_filename,
            storage_dir=storage_dir,
            max_docx_part_size_bytes=_MAX_DOCX_PART_SIZE,
        )
        await session.commit()
    return result


async def _fetch_document_or_none(source_id: uuid.UUID, file_name: str) -> Document | None:
    async with db_session() as session:
        return (
            await session.execute(
                sa.select(Document).where(
                    Document.source_id == source_id, Document.file_name == file_name
                )
            )
        ).scalar_one_or_none()


async def _fetch_document(source_id: uuid.UUID, file_name: str) -> Document:
    document = await _fetch_document_or_none(source_id, file_name)
    assert document is not None
    return document


# --- (a) each real format ingests successfully --------------------------


@pytest.mark.parametrize(
    "load_bytes, filename",
    [
        (lambda: read_pdf_fixture("multi_page.pdf"), "report.pdf"),
        (lambda: read_docx_fixture("nested_headings.docx"), "report.docx"),
        (lambda: read_text_fixture("sample.txt").encode("utf-8"), "notes.txt"),
        (lambda: read_markdown_fixture("nested_headings.md").encode("utf-8"), "notes.md"),
    ],
)
async def test_each_real_format_ingests_successfully(
    monkeypatch, live_test_services, tmp_path, load_bytes, filename
):
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)
    data = load_bytes()
    upload_id = upload_storage.save(data, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE)

    result = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename=filename,
        storage_dir=tmp_path,
    )

    assert result.status == "ingested"
    assert result.document_id is not None

    document = await _fetch_document(ids["source_id"], filename)
    assert document.id == result.document_id
    assert document.file_name == filename
    assert document.url is None  # confirms the (source_id, file_name) keying, never url
    assert document.status == "extracted"

    records, _ = await client.scroll(
        collection_name=collection_name,
        scroll_filter=qdrant.tenant_filter(ids["tenant_id"]),
        limit=10,
    )
    assert len(records) >= 1
    assert all(record.payload["source_type"] == "upload" for record in records)
    assert all(record.payload["file_name"] == filename for record in records)
    assert all(record.payload["source_url"] is None for record in records)
    assert all(record.payload["doc_id"] == str(result.document_id) for record in records)


# --- (b) identical content: skip re-chunk/re-embed, proven by call count ---


async def test_rerun_with_identical_content_skips_rechunk_and_reembed(
    monkeypatch, live_test_services, tmp_path
):
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)
    data = read_pdf_fixture("multi_page.pdf")
    upload_id = upload_storage.save(data, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE)

    first = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename="report.pdf",
        storage_dir=tmp_path,
    )
    assert first.status == "ingested"
    count_after_first = await client.count(collection_name=collection_name)

    calls = []
    original = upload_adapter_module.embed_and_upsert

    async def _counting(*args, **kwargs):
        calls.append(1)
        return await original(*args, **kwargs)

    monkeypatch.setattr(upload_adapter_module, "embed_and_upsert", _counting)

    # Re-run against the SAME stored bytes (same upload_id) -- the
    # realistic shape of "nothing changed since the last run."
    second = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename="report.pdf",
        storage_dir=tmp_path,
    )

    assert second.status == "unchanged"
    assert second.document_id == first.document_id
    assert calls == []  # embed_and_upsert() was never called the second time
    count_after_second = await client.count(collection_name=collection_name)
    assert count_after_second.count == count_after_first.count


# --- (c) changed content (a new upload_id, same filename/source): re-embeds ---


async def test_rerun_with_changed_content_reembeds_same_document(
    monkeypatch, live_test_services, tmp_path
):
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)

    first_upload_id = upload_storage.save(
        read_text_fixture("sample.txt").encode("utf-8"),
        storage_dir=tmp_path,
        max_size_bytes=_MAX_UPLOAD_SIZE,
    )
    first = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=first_upload_id,
        original_filename="notes.txt",
        storage_dir=tmp_path,
    )
    assert first.status == "ingested"
    hash_after_first = (await _fetch_document(ids["source_id"], "notes.txt")).content_hash

    # A genuinely different upload_id -- the realistic shape of a tenant
    # re-uploading changed content under the same source/filename (the
    # caller, 2.7.d's own job handler, is responsible for knowing the
    # CURRENT upload_id to pass each time; see this module's own header
    # comment on where that mapping is expected to live).
    changed_bytes = b"Completely different content, " * 20
    second_upload_id = upload_storage.save(
        changed_bytes, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE
    )
    second = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=second_upload_id,
        original_filename="notes.txt",
        storage_dir=tmp_path,
    )

    assert second.status == "ingested"
    assert second.document_id == first.document_id  # same document_id reused

    document = await _fetch_document(ids["source_id"], "notes.txt")
    assert document.content_hash != hash_after_first  # proves the row really was updated

    records, _ = await client.scroll(
        collection_name=collection_name,
        scroll_filter=qdrant.tenant_filter(ids["tenant_id"]),
        limit=10,
    )
    assert len(records) >= 1
    assert all("Completely different" in record.payload["text"] for record in records)


# --- (d) a corrupt file: definitive failure, no partial documents row ---
#
# A real finding made live while writing these tests, not assumed from
# 2.7.b's own re-confirmation (which only re-ran 2.4.d/e/f's EXISTING
# tests, never tried these specific input shapes): truncating a whole
# file by half is NOT a reachable "failed" case through ingest_upload()
# for docx or txt/md -- sniff_file_type() already opens/decodes the SAME
# bytes to classify them, so a truncated zip or invalid-UTF-8 text is
# already caught as "unknown"/"skipped" before any extractor is ever
# called (see upload_adapter.py's own header comment for the full
# writeup). Only PDF's own truncation case is genuinely reachable this
# way, since sniff_file_type() never opens/parses PDF bytes at all, only
# checks the `%PDF-` prefix. The DOCX case below uses a DIFFERENT,
# actually-reachable corruption instead: a docx-shaped zip that passes
# sniffing (valid zip, `word/document.xml` present) but whose own content
# is invalid XML -- confirmed live to raise `lxml.etree.XMLSyntaxError`,
# which upload_adapter.py's own dispatch table was widened to catch.


def _truncated_pdf() -> bytes:
    valid_bytes = read_pdf_fixture("multi_page.pdf")
    return valid_bytes[: len(valid_bytes) // 2]


def _docx_with_invalid_xml_content() -> bytes:
    # A REAL, complete .docx (built via python-docx itself, so every OPC
    # part sniff_file_type()/Document() expect is genuinely present),
    # with its own word/document.xml member's CONTENT replaced by
    # garbage -- not a truncated/broken ZIP CONTAINER (sniff_file_type()
    # would catch that first; see this section's own header comment).
    import docx as docx_module

    document = docx_module.Document()
    document.add_paragraph("hello world")
    buffer = io.BytesIO()
    document.save(buffer)

    with zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["word/document.xml"] = b"<not valid xml at all <<<"

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return out.getvalue()


@pytest.mark.parametrize(
    "corrupt_bytes, filename",
    [
        (_truncated_pdf, "broken.pdf"),
        (_docx_with_invalid_xml_content, "broken.docx"),
    ],
)
async def test_a_corrupt_file_is_a_definitive_failure_with_no_partial_document_row(
    monkeypatch, live_test_services, tmp_path, corrupt_bytes, filename
):
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)
    data = corrupt_bytes()
    upload_id = upload_storage.save(data, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE)

    result = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename=filename,
        storage_dir=tmp_path,
    )

    assert result.status == "failed"
    assert result.document_id is None
    assert result.reason is not None
    assert await _fetch_document_or_none(ids["source_id"], filename) is None


async def test_invalid_utf8_bytes_are_skipped_not_failed_since_sniffing_catches_them_first(
    monkeypatch, live_test_services, tmp_path
):
    # The other half of this section's own header-comment finding: an
    # invalid-UTF-8 "text" upload never reaches extract_text()/extract_
    # markdown() at all (sniff_file_type() already decodes-to-confirm
    # UTF-8-ness, 2.7.a) -- so its real, observed outcome through this
    # pipeline is "skipped", not "failed". Proven directly, not assumed:
    # extract_text()/extract_markdown() are both patched to raise if
    # called at all.
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("an extractor was called for invalid-UTF-8 content")

    monkeypatch.setattr(upload_adapter_module, "extract_text", _must_not_be_called)
    monkeypatch.setattr(upload_adapter_module, "extract_markdown", _must_not_be_called)

    data = b"caf\xe9 with an invalid utf-8 byte"
    upload_id = upload_storage.save(data, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE)

    result = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename="broken.txt",
        storage_dir=tmp_path,
    )

    assert result.status == "skipped"
    assert await _fetch_document_or_none(ids["source_id"], "broken.txt") is None


# --- (e) an unrecognized type: skipped, no extractor ever attempted -----


async def test_unrecognized_file_type_is_skipped_with_no_extractor_attempted(
    monkeypatch, live_test_services, tmp_path
):
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("an extractor was called for an unrecognized file type")

    monkeypatch.setattr(upload_adapter_module, "extract_pdf", _must_not_be_called)
    monkeypatch.setattr(upload_adapter_module, "extract_docx", _must_not_be_called)
    monkeypatch.setattr(upload_adapter_module, "extract_text", _must_not_be_called)
    monkeypatch.setattr(upload_adapter_module, "extract_markdown", _must_not_be_called)

    # PNG's own magic bytes -- not pdf, not zip, not valid UTF-8 either;
    # matches test_upload_sniff.py's own identical "unrecognized" probe.
    data = b"\x89PNG\r\n\x1a\n" + b"some binary payload"
    upload_id = upload_storage.save(data, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE)

    result = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename="mystery.bin",
        storage_dir=tmp_path,
    )

    assert result.status == "skipped"
    assert result.document_id is None
    assert await _fetch_document_or_none(ids["source_id"], "mystery.bin") is None


# --- (f) the zip-bomb-shaped DOCX from 2.7.b, rejected end-to-end -------


async def test_zip_bomb_shaped_docx_is_rejected_through_the_real_end_to_end_path(
    monkeypatch, live_test_services, tmp_path
):
    # The exact PoC from the Step 2.7.b decision log, run through THIS
    # module's own real orchestration, not just sniff_file_type() in
    # isolation (test_upload_sniff.py already proves the sniffing-level
    # mechanism; this proves ingest_upload() itself never reaches an
    # extractor for a bomb-shaped upload either).
    client, collection_name, ids = await _setup(monkeypatch, live_test_services)

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("extract_docx() was called for a zip-bomb-shaped upload")

    monkeypatch.setattr(upload_adapter_module, "extract_docx", _must_not_be_called)

    buffer = io.BytesIO()
    bomb = b"\x00" * (200 * 1024 * 1024)
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        archive.writestr("word/document.xml", bomb)
    data = buffer.getvalue()
    assert len(data) < 1024 * 1024  # really is bomb-shaped: tiny on disk, huge if decompressed

    upload_id = upload_storage.save(data, storage_dir=tmp_path, max_size_bytes=_MAX_UPLOAD_SIZE)

    result = await _ingest(
        client,
        collection_name,
        tenant_id=ids["tenant_id"],
        source_id=ids["source_id"],
        upload_id=upload_id,
        original_filename="totally_a_docx.docx",
        storage_dir=tmp_path,
    )

    assert result.status == "skipped"
    assert await _fetch_document_or_none(ids["source_id"], "totally_a_docx.docx") is None
